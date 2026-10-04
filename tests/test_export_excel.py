from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from fund_analyzer.export_excel import ANALYSIS_COLUMNS, export_to_excel


def test_export_excel_contains_required_sheets_columns_and_basic_formatting(tmp_path):
    output_path = Path(tmp_path) / "analysis.xlsx"
    funds = pd.DataFrame([{"fund_code": "AAA", "fund_name": "Fund", "fund_type": "Equity"}])
    prices = pd.DataFrame(
        [{"fund_code": "AAA", "date": pd.Timestamp("2026-01-01"), "price": 100.0}]
    )
    analysis_row = {column: "" for column in ANALYSIS_COLUMNS}
    analysis_row.update(
        {
            "fund_code": "AAA",
            "fund_name": "Fund",
            "fund_type": "Equity",
            "latest_price": 100.0,
            "daily_return_pct": 1.25,
            "atr_like_pct": 0.0125,
            "score": 7,
            "label": "HOLD",
        }
    )
    analysis = pd.DataFrame([analysis_row])
    portfolio = pd.DataFrame(
        [
            {
                "fund_code": "AAA",
                "fund_name": "Fund",
                "buy_date": pd.Timestamp("2026-01-01"),
                "buy_price": 100.0,
                "units": 1,
                "latest_price": 100.0,
                "current_value": 100.0,
                "cost_value": 100.0,
                "profit_loss_amount": 0.0,
                "profit_loss_pct": 0.0,
                "atr_like": 1.0,
                "atr_like_pct": 0.01,
                "stop_loss": 98.5,
                "one_atr_level": 99.0,
                "two_atr_level": 98.0,
                "three_atr_level": 97.0,
                "four_atr_level": 96.0,
                "latest_label": "HOLD",
            }
        ]
    )
    sector_summary = pd.DataFrame(
        [
            {
                "fund_type": "Equity",
                "STRONG BUY": 0,
                "EARLY BUY": 0,
                "HOLD": 1,
                "WAIT": 0,
                "SELL": 0,
                "HOLD RSI HIGH": 0,
                "TOTAL": 1,
                "net_balance": -1,
            }
        ]
    )

    export_to_excel(funds, prices, analysis, portfolio, sector_summary, output_path)
    workbook = load_workbook(output_path)

    assert workbook.sheetnames == ["Funds", "Prices", "Analysis Table", "Portfolio", "Sector Summary"]
    analysis_sheet = workbook["Analysis Table"]
    assert analysis_sheet.freeze_panes == "A2"
    assert analysis_sheet.auto_filter.ref is not None
    assert [cell.value for cell in analysis_sheet[1]] == ANALYSIS_COLUMNS
    assert analysis_sheet["N2"].number_format == '0.00"%"'
    assert analysis_sheet["AB2"].number_format == "0.00%"

