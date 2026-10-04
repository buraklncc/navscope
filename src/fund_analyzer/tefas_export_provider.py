"""Manual TEFAS CSV/XLSX export parsing and combined universe workflows."""

from __future__ import annotations

import math
import numbers
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from fund_analyzer import config
from fund_analyzer.data_provider import PORTFOLIO_COLUMNS, clean_fund_data
from fund_analyzer.models import LoadedData


TEFAS_EXPORT_COLUMNS = [
    "fund_code",
    "fund_name",
    "fund_type",
    "risk_value",
    "risk_scale",
    "return_1m_pct",
    "return_3m_pct",
    "return_6m_pct",
    "return_ytd_pct",
    "return_1y_pct",
    "return_3y_pct",
    "return_5y_pct",
]

UNIVERSE_OUTPUT_COLUMNS = [
    "fund_code",
    "fund_name",
    "fund_type",
    "risk_value",
    "return_1m_pct",
    "return_3m_pct",
    "return_6m_pct",
    "return_ytd_pct",
    "return_1y_pct",
    "return_3y_pct",
    "return_5y_pct",
]

RAW_COLUMN_MAP = {
    "fon kodu": "fund_code",
    "fonkodu": "fund_code",
    "kod": "fund_code",
    "fon adi": "fund_name",
    "fon adı": "fund_name",
    "fonun adi": "fund_name",
    "fonun adı": "fund_name",
    "fon unvani": "fund_name",
    "fon unvanı": "fund_name",
    "semsiye fon turu": "fund_type",
    "semsiye fon türü": "fund_type",
    "fon turu": "fund_type",
    "fon türü": "fund_type",
    "fonun risk degeri": "risk_raw",
    "fonun risk değeri": "risk_raw",
    "risk degeri": "risk_raw",
    "risk değeri": "risk_raw",
    "1 ay": "return_1m_pct",
    "1 ay %": "return_1m_pct",
    "1 ay (%)": "return_1m_pct",
    "3 ay": "return_3m_pct",
    "3 ay %": "return_3m_pct",
    "3 ay (%)": "return_3m_pct",
    "6 ay": "return_6m_pct",
    "6 ay %": "return_6m_pct",
    "6 ay (%)": "return_6m_pct",
    "yilbasindan itibaren": "return_ytd_pct",
    "yılbaşından itibaren": "return_ytd_pct",
    "yilbasindan itibaren %": "return_ytd_pct",
    "yılbaşından itibaren (%)": "return_ytd_pct",
    "yilbasindan itibaren (%)": "return_ytd_pct",
    "1 yil": "return_1y_pct",
    "1 yıl": "return_1y_pct",
    "1 yil %": "return_1y_pct",
    "1 yıl (%)": "return_1y_pct",
    "3 yil": "return_3y_pct",
    "3 yıl": "return_3y_pct",
    "3 yil %": "return_3y_pct",
    "3 yıl (%)": "return_3y_pct",
    "5 yil": "return_5y_pct",
    "5 yıl": "return_5y_pct",
    "5 yil %": "return_5y_pct",
    "5 yıl (%)": "return_5y_pct",
}


@dataclass
class TefasExportDataProvider:
    """Read a manually downloaded TEFAS table export as a universe-only provider."""

    export_file: Path
    write_universe_path: Path | None = None

    def load_all(self) -> LoadedData:
        warnings: list[str] = []
        funds = parse_tefas_export(self.export_file, warnings=warnings)
        if self.write_universe_path is not None:
            write_generated_universe(funds, self.write_universe_path)

        metadata = export_metadata(
            provider="tefas-export",
            export_file=self.export_file,
            export_rows_loaded=len(funds),
            export_funds_parsed=len(funds),
            funds_after_filtering=len(funds),
            write_universe_path=self.write_universe_path,
        )
        metadata.update(
            {
                "portfolio_file_used": "None",
                "portfolio_rows_analyzed": 0,
                "portfolio_skipped": True,
                "unknown_fund_types": funds_with_unknown_type(funds),
                "historical_prices_available": False,
                "requested_historical_funds": 0,
                "successfully_fetched_historical_funds": 0,
                "failed_historical_fund_codes": [],
            }
        )
        warnings.append(
            "TEFAS export mode contains universe and ready-made return data only; technical indicators were not calculated."
        )
        return LoadedData(
            funds=funds,
            prices=pd.DataFrame(columns=["fund_code", "date", "price"]),
            portfolio=pd.DataFrame(columns=PORTFOLIO_COLUMNS),
            warnings=warnings,
            metadata=metadata,
        )


def parse_tefas_export(path: Path | str, warnings: list[str] | None = None) -> pd.DataFrame:
    """Parse a TEFAS table export from CSV or Excel into normalized columns."""

    path = Path(path)
    raw = _read_export_file(path)
    normalized = normalize_tefas_export_frame(raw)
    if normalized.empty and warnings is not None:
        warnings.append(f"{path}: no fund rows could be parsed from TEFAS export.")
    return normalized


def extract_tefas_export_metadata(path: Path | str) -> dict[str, object]:
    """Read TEFAS report-header metadata without changing the table contract.

    The real report time lives above the table header in Excel exports.  It is
    intentionally kept separate from the time a user uploads the file.
    """

    path = Path(path)
    metadata: dict[str, object] = {
        "export_generated_at": None,
        "reported_record_count": None,
    }
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _extract_xlsx_export_metadata(path, metadata)
    if path.suffix.lower() == ".csv":
        return _extract_csv_export_metadata(path, metadata)
    return metadata


def normalize_tefas_export_frame(raw: pd.DataFrame) -> pd.DataFrame:
    """Normalize raw export rows and Turkish headers."""

    data = _promote_header_if_needed(raw)
    excel_percent_targets = _excel_percent_return_targets(data)
    column_map = {_original: _canonical_column(_original) for _original in data.columns}
    data = data.rename(columns={source: target for source, target in column_map.items() if target})
    keep_columns = [column for column in data.columns if column in {*TEFAS_EXPORT_COLUMNS, "risk_raw"}]
    data = data[keep_columns].copy()

    for column in TEFAS_EXPORT_COLUMNS:
        if column not in data.columns:
            data[column] = pd.NA

    data["fund_code"] = data["fund_code"].astype("string").str.strip().str.upper()
    data["fund_name"] = data["fund_name"].astype("string").str.strip()
    data["fund_type"] = data["fund_type"].astype("string").str.strip()
    data.loc[data["fund_type"].isna() | (data["fund_type"] == ""), "fund_type"] = config.TEFAS_UNKNOWN_FUND_TYPE

    risk_source = data["risk_raw"] if "risk_raw" in data.columns else data["risk_value"]
    risk_parts = risk_source.apply(parse_risk_value)
    data["risk_value"] = risk_parts.apply(lambda item: item[0])
    data["risk_scale"] = risk_parts.apply(lambda item: item[1])

    for column in _return_columns():
        data[column] = data[column].apply(
            lambda value, column=column: parse_turkish_percent(
                value,
                numeric_fraction=column in excel_percent_targets,
            )
        )

    data = data.loc[_valid_fund_code_mask(data["fund_code"]) & _valid_text_mask(data["fund_name"])].copy()
    data = data[TEFAS_EXPORT_COLUMNS].drop_duplicates("fund_code", keep="last").reset_index(drop=True)
    for column in ["risk_value", "risk_scale", *_return_columns()]:
        data[column] = pd.to_numeric(data[column], errors="coerce")
    return data[TEFAS_EXPORT_COLUMNS]


def filter_export_universe(
    universe: pd.DataFrame,
    fund_type: str | None = None,
    risk_min: int | None = None,
    risk_max: int | None = None,
    min_return_1m: float | None = None,
    min_return_3m: float | None = None,
    min_return_6m: float | None = None,
    min_return_ytd: float | None = None,
    min_return_1y: float | None = None,
    include_funds: list[str] | None = None,
    exclude_funds: list[str] | None = None,
    limit: int | None = None,
    sort_by: str | None = None,
    sort_desc: bool = False,
    sort_asc: bool = False,
) -> pd.DataFrame:
    """Filter and order an export universe for combined historical fetching."""

    filtered = universe.copy()
    if fund_type:
        wanted = _normalize_text(fund_type)
        filtered = filtered.loc[filtered["fund_type"].apply(_normalize_text) == wanted].copy()

    if risk_min is not None:
        filtered = filtered.loc[filtered["risk_value"].fillna(math.inf) >= risk_min].copy()
    if risk_max is not None:
        filtered = filtered.loc[filtered["risk_value"].fillna(-math.inf) <= risk_max].copy()

    return_filters = {
        "return_1m_pct": min_return_1m,
        "return_3m_pct": min_return_3m,
        "return_6m_pct": min_return_6m,
        "return_ytd_pct": min_return_ytd,
        "return_1y_pct": min_return_1y,
    }
    for column, threshold in return_filters.items():
        if threshold is not None:
            filtered = filtered.loc[filtered[column].fillna(-math.inf) >= threshold].copy()

    include_codes = _code_list(include_funds)
    exclude_codes = set(_code_list(exclude_funds))

    filtered = _sort_export_universe(filtered, sort_by=sort_by, sort_desc=sort_desc, sort_asc=sort_asc)
    if limit is not None:
        filtered = filtered.head(limit).copy()

    if include_codes:
        additions = universe.loc[universe["fund_code"].isin(include_codes)]
        filtered = pd.concat([filtered, additions], ignore_index=True).drop_duplicates("fund_code", keep="first")

    if exclude_codes:
        filtered = filtered.loc[~filtered["fund_code"].isin(exclude_codes)].copy()

    return filtered.reset_index(drop=True)


def enrich_with_export_metadata(analysis: pd.DataFrame, export_universe: pd.DataFrame) -> pd.DataFrame:
    """Attach export metadata columns to an analysis table."""

    export_columns = [column for column in TEFAS_EXPORT_COLUMNS if column != "fund_name"]
    metadata = export_universe[["fund_code", *[column for column in export_columns if column != "fund_code"]]].copy()
    overlap = [column for column in metadata.columns if column in analysis.columns and column != "fund_code"]
    analysis_without_overlap = analysis.drop(columns=overlap)
    return analysis_without_overlap.merge(metadata, on="fund_code", how="left")


def write_generated_universe(universe: pd.DataFrame, output_path: Path | str) -> Path:
    """Write a clean universe CSV generated from a TEFAS export."""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    universe[UNIVERSE_OUTPUT_COLUMNS].to_csv(output_path, index=False)
    return output_path


def export_metadata(
    provider: str,
    export_file: Path | str | None,
    export_rows_loaded: int,
    export_funds_parsed: int,
    funds_after_filtering: int,
    write_universe_path: Path | str | None = None,
) -> dict[str, object]:
    return {
        "provider": provider,
        "export_file_used": str(export_file) if export_file else "None",
        "export_rows_loaded": export_rows_loaded,
        "export_funds_parsed": export_funds_parsed,
        "funds_after_filtering": funds_after_filtering,
        "write_universe_path": str(write_universe_path) if write_universe_path else "None",
    }


def funds_with_unknown_type(funds: pd.DataFrame) -> list[str]:
    if funds.empty or "fund_type" not in funds:
        return []
    unknown = funds["fund_type"].isna() | (funds["fund_type"].astype(str).str.strip() == "") | (
        funds["fund_type"].astype(str).str.strip().str.lower() == config.TEFAS_UNKNOWN_FUND_TYPE.lower()
    )
    return sorted(funds.loc[unknown, "fund_code"].tolist())


def parse_risk_value(value: Any) -> tuple[float | None, float | None]:
    if _is_empty_value(value):
        return (None, None)
    text = str(value).strip()
    match = re.search(r"(\d+(?:[,.]\d+)?)\s*/\s*(\d+(?:[,.]\d+)?)", text)
    if match:
        return (_number_or_none(match.group(1)), _number_or_none(match.group(2)))
    return (_number_or_none(text), None)


def parse_turkish_percent(value: Any, numeric_fraction: bool = False) -> float | None:
    if _is_empty_value(value):
        return None
    if isinstance(value, numbers.Real) and not isinstance(value, bool):
        parsed = float(value)
        return parsed * 100 if numeric_fraction else parsed
    text = str(value).strip().replace("%", "").replace("\u00a0", " ")
    text = text.replace(" ", "")
    return _number_or_none(text)


def _read_export_file(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _read_xlsx_with_format_metadata(path)
    if suffix == ".xls":
        return pd.read_excel(path, dtype=object, header=0)
    if suffix == ".csv":
        return _read_csv_with_fallbacks(path)
    raise ValueError(f"Unsupported TEFAS export file type: {path.suffix}")


def _read_csv_with_fallbacks(path: Path) -> pd.DataFrame:
    last_error: Exception | None = None
    for encoding in ["utf-8-sig", "utf-8", "cp1254", "iso-8859-9"]:
        try:
            return pd.read_csv(path, dtype=object, sep=None, engine="python", encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
    if last_error:
        raise last_error
    return pd.read_csv(path, dtype=object, sep=None, engine="python")


def _read_xlsx_with_format_metadata(path: Path) -> pd.DataFrame:
    raw = pd.read_excel(path, dtype=object, header=0)
    raw.attrs["excel_percent_columns_by_position"] = _xlsx_percent_columns_by_position(path)
    return raw


def _extract_xlsx_export_metadata(path: Path, metadata: dict[str, object]) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        for row in worksheet.iter_rows(max_row=20, values_only=True):
            values = list(row)
            for index, value in enumerate(values):
                _capture_report_metadata(metadata, value, values[index + 1 :])
    finally:
        workbook.close()
    return metadata


def _extract_csv_export_metadata(path: Path, metadata: dict[str, object]) -> dict[str, object]:
    lines: list[str] = []
    for encoding in ["utf-8-sig", "utf-8", "cp1254", "iso-8859-9"]:
        try:
            lines = path.read_text(encoding=encoding).splitlines()
            break
        except UnicodeDecodeError:
            continue
    for line in lines[:20]:
        values = [value.strip() for value in re.split(r"[;,]", line)]
        for index, value in enumerate(values):
            _capture_report_metadata(metadata, value, values[index + 1 :])
    return metadata


def _capture_report_metadata(metadata: dict[str, object], label_value: Any, remaining_values: list[Any]) -> None:
    if label_value is None:
        return
    raw_label = str(label_value).strip()
    label = _normalize_text(raw_label)
    candidate = next((value for value in remaining_values if value not in {None, ""}), None)
    if "disa aktarim tarihi" in label:
        value = candidate or raw_label.split(":", 1)[-1].strip()
        parsed = _parse_export_datetime(value)
        if parsed:
            metadata["export_generated_at"] = parsed
    if "toplam kayit sayisi" in label or "toplam kayit" in label:
        value = candidate or raw_label.split(":", 1)[-1].strip()
        parsed_count = _number_or_none(value)
        if parsed_count is not None:
            metadata["reported_record_count"] = int(parsed_count)


def _parse_export_datetime(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat()
    text = str(value).strip()
    for pattern in (
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
    ):
        try:
            return datetime.strptime(text, pattern).isoformat()
        except ValueError:
            continue
    return None


def _xlsx_percent_columns_by_position(path: Path) -> set[int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    percent_columns: set[int] = set()
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None and "%" in str(cell.number_format):
                percent_columns.add(cell.column - 1)
    workbook.close()
    return percent_columns


def _promote_header_if_needed(raw: pd.DataFrame) -> pd.DataFrame:
    if _has_required_columns(raw.columns):
        promoted = raw.dropna(how="all").reset_index(drop=True)
        promoted.attrs.update(raw.attrs)
        return promoted

    for row_index, row in raw.head(15).iterrows():
        values = ["" if pd.isna(value) else str(value) for value in row.tolist()]
        if _has_required_columns(values):
            promoted = raw.iloc[row_index + 1 :].copy()
            promoted.columns = values
            promoted = promoted.dropna(how="all").reset_index(drop=True)
            promoted.attrs.update(raw.attrs)
            return promoted
    promoted = raw.dropna(how="all").reset_index(drop=True)
    promoted.attrs.update(raw.attrs)
    return promoted


def _excel_percent_return_targets(data: pd.DataFrame) -> set[str]:
    percent_positions = data.attrs.get("excel_percent_columns_by_position", set())
    if not percent_positions:
        return set()
    targets: set[str] = set()
    for position, column in enumerate(data.columns):
        canonical = _canonical_column(column)
        if position in percent_positions and canonical in _return_columns():
            targets.add(canonical)
    return targets



def _has_required_columns(columns: Any) -> bool:
    canonical = {_canonical_column(column) for column in columns}
    return {"fund_code", "fund_name"}.issubset(canonical)


def _canonical_column(column: Any) -> str | None:
    normalized = _normalize_text(str(column))
    return RAW_COLUMN_MAP.get(normalized)


def _valid_fund_code_mask(codes: pd.Series) -> pd.Series:
    return codes.notna() & codes.astype(str).str.match(r"^[A-Z0-9]{2,10}$")


def _valid_text_mask(values: pd.Series) -> pd.Series:
    return values.notna() & (values.astype(str).str.strip() != "")


def _sort_export_universe(
    universe: pd.DataFrame,
    sort_by: str | None = None,
    sort_desc: bool = False,
    sort_asc: bool = False,
) -> pd.DataFrame:
    if sort_by and sort_by in universe.columns:
        if sort_desc:
            ascending = False
        elif sort_asc:
            ascending = True
        else:
            ascending = not sort_by.startswith("return")
        return universe.sort_values([sort_by, "fund_code"], ascending=[ascending, True])

    if "return_1m_pct" in universe and universe["return_1m_pct"].notna().any():
        return universe.sort_values(["return_1m_pct", "fund_code"], ascending=[False, True])
    return universe.sort_values("fund_code")


def _return_columns() -> list[str]:
    return [
        "return_1m_pct",
        "return_3m_pct",
        "return_6m_pct",
        "return_ytd_pct",
        "return_1y_pct",
        "return_3y_pct",
        "return_5y_pct",
    ]


def _code_list(values: list[str] | None) -> list[str]:
    if not values:
        return []
    return [str(value).strip().upper() for value in values if str(value).strip()]


def _normalize_text(value: Any) -> str:
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("ı", "i")
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _number_or_none(value: Any) -> float | None:
    if _is_empty_value(value):
        return None
    text = str(value).strip()
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _is_empty_value(value: Any) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except TypeError:
        pass
    return str(value).strip() in {"", "-", "–", "—"}

