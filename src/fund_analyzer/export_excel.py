"""Excel export helpers."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


ANALYSIS_COLUMNS = [
    "fund_code",
    "fund_name",
    "fund_type",
    "risk_value",
    "risk_scale",
    "return_1m_pct",
    "return_3m_pct",
    "return_6m_pct",
    "return_ytd_pct",
    "return_1y_pct",
    "return_3y_pct",
    "return_5y_pct",
    "latest_price",
    "daily_return_pct",
    "weekly_return_pct",
    "monthly_return_pct",
    "ema10",
    "ema21",
    "ema10_above_ema21",
    "ema10_crossed_above_ema21",
    "macd",
    "macd_signal",
    "macd_histogram",
    "histogram_trend_text",
    "rsi_14",
    "rsi_zone",
    "atr_like",
    "atr_like_pct",
    "score",
    "label",
    "explanation",
    "category",
    "category_name",
    "category_confidence",
    "six_month_return_pct",
    "twelve_month_return_pct",
    "ytd_return_pct",
    "annualized_return",
    "annualized_volatility",
    "downside_deviation",
    "sharpe",
    "sortino",
    "calmar",
    "max_drawdown",
    "current_drawdown",
    "average_drawdown",
    "recovery_time_days",
    "worst_day",
    "worst_week",
    "worst_month",
    "positive_day_ratio",
    "positive_month_ratio",
    "category_rank_return",
    "category_rank_risk_adjusted",
    "category_rank_drawdown",
    "category_percentile",
    "momentum_score",
    "risk_score",
    "consistency_score",
    "drawdown_score",
    "data_quality_score",
    "data_quality",
    "risk_level",
    "confidence",
    "short_explanation",
    "strategy_version",
]

PERCENT_POINT_COLUMNS = {
    "daily_return_pct",
    "weekly_return_pct",
    "monthly_return_pct",
    "three_month_return_pct",
    "profit_loss_pct",
    "return_1m_pct",
    "return_3m_pct",
    "return_6m_pct",
    "return_ytd_pct",
    "return_1y_pct",
    "return_3y_pct",
    "return_5y_pct",
    "six_month_return_pct",
    "twelve_month_return_pct",
    "ytd_return_pct",
}

FRACTION_PERCENT_COLUMNS = {
    "atr_like_pct",
    "annualized_return",
    "annualized_volatility",
    "downside_deviation",
    "max_drawdown",
    "current_drawdown",
    "average_drawdown",
    "worst_day",
    "worst_week",
    "worst_month",
    "positive_day_ratio",
    "positive_month_ratio",
}
TWO_DECIMAL_COLUMNS = {
    "price",
    "latest_price",
    "ema10",
    "ema21",
    "macd",
    "macd_signal",
    "macd_histogram",
    "rsi_14",
    "rsi_value",
    "atr_like",
    "buy_price",
    "units",
    "current_value",
    "cost_value",
    "profit_loss_amount",
    "stop_loss",
    "one_atr_level",
    "two_atr_level",
    "three_atr_level",
    "four_atr_level",
}


def export_to_excel(
    funds: pd.DataFrame,
    prices: pd.DataFrame,
    analysis: pd.DataFrame,
    portfolio: pd.DataFrame,
    sector_summary: pd.DataFrame,
    output_path: Path,
    batch_errors: pd.DataFrame | None = None,
) -> Path:
    """Export all result tables to a multi-sheet Excel workbook."""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        funds.to_excel(writer, sheet_name="Funds", index=False)
        prices.to_excel(writer, sheet_name="Prices", index=False)
        analysis[_analysis_export_columns(analysis)].to_excel(writer, sheet_name="Analysis Table", index=False)
        portfolio.to_excel(writer, sheet_name="Portfolio", index=False)
        sector_summary.to_excel(writer, sheet_name="Sector Summary", index=False)
        if batch_errors is not None and not batch_errors.empty:
            batch_errors.to_excel(writer, sheet_name="Batch Errors", index=False)
        _format_workbook(writer)

    return output_path


def _format_workbook(writer: pd.ExcelWriter) -> None:
    workbook = writer.book
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    label_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        header_by_column = {cell.column: str(cell.value) for cell in worksheet[1] if cell.value}

        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for column_index, header in header_by_column.items():
            number_format = _number_format_for_header(header)
            for cell in worksheet.iter_rows(
                min_row=2,
                max_row=worksheet.max_row,
                min_col=column_index,
                max_col=column_index,
            ):
                data_cell = cell[0]
                if number_format and data_cell.value is not None:
                    data_cell.number_format = number_format
                if header in {"score", "label", "latest_label"}:
                    data_cell.fill = label_fill
                    data_cell.font = Font(bold=True)

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)


def _number_format_for_header(header: str) -> str | None:
    if header in PERCENT_POINT_COLUMNS:
        return '0.00"%"'
    if header in FRACTION_PERCENT_COLUMNS:
        return "0.00%"
    if header in TWO_DECIMAL_COLUMNS:
        return "0.00"
    return None


def _analysis_export_columns(analysis: pd.DataFrame) -> list[str]:
    return [column for column in ANALYSIS_COLUMNS if column in analysis.columns]

